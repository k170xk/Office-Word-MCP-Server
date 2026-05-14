import pytest
from docx import Document

from word_document_server.tools.document_tools import create_document
from word_document_server.tools import content_tools


@pytest.mark.asyncio
async def test_add_table_string_data_rejected(tmp_path):
    doc_path = tmp_path / "t.docx"
    await create_document(str(doc_path), title="Test")
    out = await content_tools.add_table(str(doc_path), rows=3, cols=3, data="abc")
    assert "Invalid `data`" in out


@pytest.mark.asyncio
async def test_add_table_nested_data_fills_cells(tmp_path):
    doc_path = tmp_path / "t.docx"
    await create_document(str(doc_path))
    rows = [["A", "B", "C"], ["1", "2", "3"]]
    await content_tools.add_table(str(doc_path), rows=2, cols=3, data=rows)
    doc = Document(str(doc_path))
    t = doc.tables[-1]
    assert len(t.rows) >= 2
    assert t.cell(0, 0).text == "A"
    assert t.cell(0, 2).text == "C"
    assert t.cell(1, 1).text == "2"


@pytest.mark.asyncio
async def test_add_table_flat_scalar_list_single_column(tmp_path):
    doc_path = tmp_path / "t.docx"
    await create_document(str(doc_path))
    await content_tools.add_table(str(doc_path), rows=0, cols=2, data=["x", "y"])
    doc = Document(str(doc_path))
    t = doc.tables[-1]
    assert len(t.rows) == 2
    assert t.cell(0, 0).text == "x"
    assert t.cell(0, 1).text == ""
    assert t.cell(1, 0).text == "y"


@pytest.mark.asyncio
async def test_append_table_rows_string_rejected(tmp_path):
    doc_path = tmp_path / "t.docx"
    await create_document(str(doc_path))
    await content_tools.add_table(str(doc_path), rows=1, cols=2, data=[["H1", "H2"]])
    out = await content_tools.append_table_rows(str(doc_path), 0, "bad")
    assert "Invalid `data`" in out


@pytest.mark.asyncio
async def test_append_table_rows_nested_ok(tmp_path):
    doc_path = tmp_path / "t.docx"
    await create_document(str(doc_path))
    await content_tools.add_table(str(doc_path), rows=1, cols=2, data=[["H1", "H2"]])
    await content_tools.append_table_rows(str(doc_path), 0, [["a", "b"]])
    doc = Document(str(doc_path))
    t = doc.tables[0]
    assert len(t.rows) == 2
    assert t.cell(1, 0).text == "a"
    assert t.cell(1, 1).text == "b"


@pytest.mark.asyncio
async def test_add_table_header_shading_default_brand_green(tmp_path):
    doc_path = tmp_path / "t.docx"
    await create_document(str(doc_path))
    await content_tools.add_table(str(doc_path), rows=1, cols=2, data=[["A", "B"]])
    doc = Document(str(doc_path))
    cell_xml = doc.tables[-1].cell(0, 0)._tc.xml
    assert "77C343" in cell_xml


@pytest.mark.asyncio
async def test_add_table_no_header_style_when_disabled(tmp_path):
    doc_path = tmp_path / "t.docx"
    await create_document(str(doc_path))
    await content_tools.add_table(
        str(doc_path), rows=1, cols=2, data=[["A", "B"]], style_header_row=False
    )
    doc = Document(str(doc_path))
    cell_xml = doc.tables[-1].cell(0, 0)._tc.xml
    assert "77C343" not in cell_xml
